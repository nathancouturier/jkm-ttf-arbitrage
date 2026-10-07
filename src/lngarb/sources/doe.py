"""US LNG exports cargo by cargo, from the US Department of Energy.

The Department of Energy's Office of Fossil Energy and Carbon Management
publishes, with each monthly report on natural gas imports and exports, a
workbook of every LNG export and re-export since January 2016, one row per
cargo or part of a cargo: the departure date, the company holding the export
authorisation, its docket and term, the supplier, the mode of transport, the
ship, the port of exit, the destination country and the volume in MMcf. EIA's
monthly exports by country are built from the same reports, and this file
reproduces them to the MMcf; it adds what the monthly table cannot show, which
terminal a cargo left from and on which day.

The workbook's address changes every month and every year. It is read from the
year's report page, which is read from the list of report pages, and the parse
fails when either link is missing or ambiguous rather than guess. Every
destination must carry a region of lngarb.config.EIA_DESTINATIONS, by country
name, so a new destination stops the parse.

DOE's site says: "Government information at DOE websites is in the public
domain." Each workbook read is also kept in data/private/doe/.
"""

from __future__ import annotations

import io
import re
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
import pandas as pd
from bs4 import BeautifulSoup

from ..config import BOUNDS_CARGO_MMCF, EIA_DESTINATIONS
from . import base
from .base import Adapter, SourceError, http_get

__all__ = [
    "LIST_URL",
    "discover_report_page",
    "discover_workbook_url",
    "parse_cargo_workbook",
    "DoeLngExportCargoes",
]

LIST_URL = "https://www.energy.gov/hgeo/listings/natural-gas-imports-and-exports-monthly-reports"
SHEET = "By Vessel and ISO Container"
HEADER = (
    "Arrival/Departure Date", "Companies", "Docket Number", "Docket Term", "Activity",
    "Gas Type", "Supplier", "Mode of Transport", "Tanker", "Point of Entry or Exit",
    "Country", "Volume (MMCF)", "U.S. Contiguous",
)
COLUMNS = (
    "date", "company", "docket_number", "docket_term", "activity", "gas_type",
    "supplier", "mode", "tanker", "point_of_exit", "country", "volume_mmcf", "us_contiguous",
)

_REPORT_PAGE = re.compile(r"/articles/natural-gas-imports-and-exports-monthly-(\d{4})$")
_WORKBOOK_TEXT = re.compile(r"U\.S\. LNG Exports and Re-Exports Details \(Jan 2016 - ([A-Z][a-z]{2}) (\d{4})\)")


def _region_of_country() -> dict[str, str]:
    regions: dict[str, str] = {}
    for destination in EIA_DESTINATIONS.values():
        known = regions.get(destination.country)
        if known is not None and known != destination.region:
            raise SourceError("%s is mapped to two regions" % destination.country)
        regions[destination.country] = destination.region
    return regions


def discover_report_page(html: str | bytes, *, base_url: str = LIST_URL) -> str:
    """The address of the latest year's report page on the list of report pages."""
    soup = BeautifulSoup(html, "lxml")
    found: dict[int, str] = {}
    for link in soup.find_all("a", href=True):
        href = link["href"].split("?")[0].rstrip("/")
        match = _REPORT_PAGE.search(href)
        if match:
            found[int(match.group(1))] = urljoin(base_url, href)
    if not found:
        raise SourceError("the list of report pages links no year's report page")
    return found[max(found)]


def discover_workbook_url(html: str | bytes, *, base_url: str) -> tuple[str, str]:
    """The cargo workbook's address on a year's report page, and the month it runs to."""
    soup = BeautifulSoup(html, "lxml")
    found = []
    for link in soup.find_all("a", href=True):
        text = re.sub(r"\s+", " ", link.get_text()).strip()
        match = _WORKBOOK_TEXT.search(text)
        if match and link["href"].lower().split("?")[0].endswith(".xlsx"):
            found.append((urljoin(base_url, link["href"]), "%s %s" % match.groups()))
    if len(found) != 1:
        raise SourceError(
            "the report page links %d cargo workbooks where one was expected" % len(found)
        )
    return found[0]


def parse_cargo_workbook(payload: bytes) -> tuple[pd.DataFrame, int]:
    """Every cargo row of the workbook, oldest first, with its region.

    Returns the frame and the number of rows left out because they carry no
    date. Raises SourceError when the sheet or its header is not the one
    described above, a volume is not a number, or a destination has no region.
    """
    book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    if SHEET not in book.sheetnames:
        raise SourceError("the workbook has no sheet %r" % SHEET)
    rows = book[SHEET].iter_rows(values_only=True)
    header = tuple(str(c).strip() if c is not None else "" for c in next(rows))
    if header[: len(HEADER)] != HEADER:
        raise SourceError("the cargo sheet's header is %r, not the one this parser reads" % (header,))
    regions = _region_of_country()
    records, undated = [], 0
    for row in rows:
        if row is None or all(c is None for c in row[: len(HEADER)]):
            continue
        record = dict(zip(COLUMNS, row[: len(HEADER)]))
        if not isinstance(record["date"], datetime):
            undated += 1
            continue
        volume = record["volume_mmcf"]
        if not isinstance(volume, (int, float)) or isinstance(volume, bool):
            raise SourceError("a cargo of %s carries a volume of %r" % (record["date"].date(), volume))
        country = str(record["country"]).strip()
        if country not in regions:
            raise SourceError(
                "destination %r has no region in lngarb.config.EIA_DESTINATIONS" % country
            )
        record["country"] = country
        record["region"] = regions[country]
        record["date"] = pd.Timestamp(record["date"].date())
        records.append(record)
    if not records:
        raise SourceError("the cargo sheet has no dated row")
    frame = pd.DataFrame(records)
    for column in ("company", "docket_number", "docket_term", "activity", "gas_type",
                   "supplier", "mode", "tanker", "point_of_exit", "us_contiguous"):
        frame[column] = frame[column].map(lambda v: "" if v is None else re.sub(r"\s+", " ", str(v)).strip())
    frame["volume_mmcf"] = frame["volume_mmcf"].astype(float)
    ordered = list(COLUMNS[:11]) + ["region", "volume_mmcf", "us_contiguous"]
    frame = frame[ordered].sort_values(
        ["date", "point_of_exit", "tanker", "country", "company"], kind="mergesort"
    ).reset_index(drop=True)
    return frame, undated


def private_dir() -> Path:
    """Where each workbook read is kept, gitignored."""
    return base.PRIVATE / "doe"


class DoeLngExportCargoes(Adapter):
    """US LNG exports and re-exports, cargo by cargo, from January 2016."""

    name = "doe_lng_export_cargoes"
    source = "U.S. Department of Energy, Office of Fossil Energy and Carbon Management, U.S. LNG Exports and Re-Exports Details"
    url = LIST_URL
    page_url = LIST_URL
    unit = "MMcf per cargo"
    frequency = "monthly"
    method = "published"
    unique_dates = False
    required_cols = ("date", "point_of_exit", "country", "region", "volume_mmcf")
    bounds = {"volume_mmcf": BOUNDS_CARGO_MMCF}
    min_observations = {"volume_mmcf": 5000}
    observation_column = "volume_mmcf"
    min_rows = 5000

    def __init__(self, *, from_file: Path | None = None, fetched_at: str | None = None, period: str | None = None):
        self.from_file = from_file
        self.fetched_at = fetched_at
        self.period = period

    def fetch(self) -> pd.DataFrame:
        if self.from_file is not None:
            payload = Path(self.from_file).read_bytes()
            period = self.period or "unknown"
            how = "read from a copy fetched at %s" % self.fetched_at
        else:
            page = discover_report_page(http_get(LIST_URL).content)
            workbook_url, period = discover_workbook_url(http_get(page).content, base_url=page)
            payload = http_get(workbook_url, timeout=120).content
            self.url = workbook_url
            how = "fetched from %s" % workbook_url
        frame, undated = parse_cargo_workbook(payload)
        folder = private_dir()
        folder.mkdir(parents=True, exist_ok=True)
        (folder / ("US_LNG_Exports_ReExports_Details_to_%s.xlsx" % period.replace(" ", "_"))).write_bytes(payload)
        self.vintage = "January 2016 to %s" % period
        self.note = (
            "%s. One row per cargo or part of a cargo, dated by departure; %d row(s) "
            "without a date are left out. Re-exports and ISO containers are kept and "
            "labelled. Acknowledgement: U.S. Department of Energy." % (how, undated)
        )
        return frame


def main() -> int:
    try:
        entry = DoeLngExportCargoes().run()
    except Exception as exc:
        print("FAILED doe_lng_export_cargoes: %s" % exc)
        return 1
    print("ok     %s: %d cargo rows, %s to %s, %s" % (
        entry["series"], entry["file_rows"], entry["first_date"], entry["last_date"], entry["vintage"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
