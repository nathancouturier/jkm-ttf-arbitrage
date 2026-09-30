"""World Bank Pink Sheet gas prices, monthly, USD/MMBtu: Europe, the US and Japan LNG.

The workbook "CMO-Historical-Data-Monthly.xlsx" is linked from the World Bank's
commodity markets page. The document id in its path changes, so the link is
read from the page every time and the adapter fails when it cannot find exactly
one. Series are found by the label printed above them and checked by the unit
printed below the label, never by position.

The three series and what the World Bank says each one is, from the workbook's
Description sheet:

    Natural gas, Europe             "from April 2015, Netherlands Title Transfer
                                    Facility (TTF); April 2010 to March 2015,
                                    average import border price and a spot price
                                    component, including UK; during June 2000 -
                                    March 2010 prices excludes UK."
    Natural gas, US                 "spot price at Henry Hub, Louisiana"
    Liquefied natural gas, Japan    "LNG, import price, cif; recent two months'
                                    averages are estimates."

(the dash in the first description is an ASCII hyphen in the source.) The Japan
series is an import price, not a spot price and not a JKM proxy, and its last
two months are revised: June 2026 read 12.83 in the release of 2 July 2026 and
11.79 in the release of 2 September 2026. Its last two months are therefore
marked provisional, and every value a release changes goes to a revisions log.

Months from January 2015 are kept: the study's monthly history starts there and
TTF enters the Europe series in April 2015.

This module is adapted from the World Bank adapter of the sibling repository
crack-spread-study, which reads the Europe series alone.
"""

from __future__ import annotations

import io
import math
import re
import sys
from html.parser import HTMLParser
from typing import Any, Mapping
from urllib.parse import urljoin

import pandas as pd

from ..config import BOUNDS_HENRY_HUB_USD_MMBTU, BOUNDS_LNG_USD_MMBTU
from .base import Adapter, SourceError, http_get, read_cache

__all__ = [
    "LANDING_PAGE",
    "SERIES",
    "discover_workbook_url",
    "parse_workbook",
    "compare_releases",
    "WorldBankGasMonthly",
    "WorldBankGasRevisions",
]

LANDING_PAGE = "https://www.worldbank.org/en/research/commodity-markets"
WORKBOOK_FILENAME = "CMO-Historical-Data-Monthly.xlsx"
SHEET = "Monthly Prices"
EXPECTED_UNIT = "($/mmbtu)"
FIRST_MONTH = pd.Timestamp("2015-01-01")

#: column in the cache -> the label the World Bank prints above it
SERIES: Mapping[str, str] = {
    "europe_gas_usd_mmbtu": "Natural gas, Europe",
    "henry_hub_usd_mmbtu": "Natural gas, US",
    "japan_lng_import_usd_mmbtu": "Liquefied natural gas, Japan",
}

#: The column whose last two months the World Bank calls estimates.
PROVISIONAL_COLUMN = "japan_lng_import_usd_mmbtu"
PROVISIONAL_MONTHS = 2

#: Three spellings of the missing value token live in the workbook and its
#: notes: an ellipsis character, three full stops, and two.
MISSING_TOKENS = frozenset({"", "..", "...", chr(0x2026), "n/a", "na", "nan"})


class _LinkCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.hrefs: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() != "a":
            return
        for name, value in attrs:
            if name.lower() == "href" and value:
                self.hrefs.append(value.strip())


def discover_workbook_url(html: str | bytes, *, base_url: str = LANDING_PAGE) -> str:
    """The monthly workbook's URL, read from the landing page. Exactly one, or an error."""
    if isinstance(html, (bytes, bytearray)):
        html = html.decode("utf-8", errors="replace")
    parser = _LinkCollector()
    parser.feed(html)
    hits: list[str] = []
    for href in parser.hrefs:
        path = href.split("?", 1)[0].split("#", 1)[0]
        if path.lower().rsplit("/", 1)[-1] == WORKBOOK_FILENAME.lower():
            absolute = urljoin(base_url, href)
            if absolute not in hits:
                hits.append(absolute)
    if not hits:
        raise SourceError(
            "no link to %s on %s. The document id in the path changes, so there "
            "is no URL to fall back to; find where the workbook moved"
            % (WORKBOOK_FILENAME, base_url)
        )
    if len(hits) > 1:
        raise SourceError(
            "%d different links to %s on %s: %s. Picking one would be a guess"
            % (len(hits), WORKBOOK_FILENAME, base_url, ", ".join(hits))
        )
    return hits[0]


_PERIOD = re.compile(r"^\s*(\d{4})M(\d{1,2})\s*$", re.IGNORECASE)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(raw: Any, *, where: str) -> float:
    if raw is None:
        return math.nan
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return float(raw)
    text = str(raw).strip()
    if text.lower() in MISSING_TOKENS:
        return math.nan
    try:
        return float(text.replace(",", ""))
    except ValueError as exc:
        raise SourceError(
            "%s: cannot read %r as a price. A new token for a missing month is "
            "added to MISSING_TOKENS deliberately, never coerced" % (where, raw)
        ) from exc


def _find_column(rows: list[list], label: str) -> tuple[int, int]:
    wanted = label.lower()
    found = [
        (r, c)
        for r in range(min(12, len(rows)))
        for c, cell in enumerate(rows[r])
        if _text(cell).lower() == wanted
    ]
    if len(found) != 1:
        raise SourceError(
            "%d cells in the first 12 rows of %r read %r; exactly one is needed"
            % (len(found), SHEET, label)
        )
    name_row, column = found[0]
    unit = _text(rows[name_row + 1][column]).lower() if name_row + 1 < len(rows) else ""
    if unit != EXPECTED_UNIT:
        raise SourceError(
            "the unit under %r reads %r, expected %r" % (label, unit, EXPECTED_UNIT)
        )
    return name_row, column


def parse_workbook(payload: bytes) -> tuple[pd.DataFrame, str]:
    """The three gas series from January 2015, and the release's 'Updated on' line."""
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001, openpyxl raises several types
        raise SourceError(
            "the download did not open as an xlsx workbook, %s: %s"
            % (type(exc).__name__, exc)
        ) from exc
    try:
        if SHEET not in book.sheetnames:
            raise SourceError("no %r sheet; sheets are %s" % (SHEET, ", ".join(book.sheetnames)))
        rows = [list(row) for row in book[SHEET].iter_rows(values_only=True)]
    finally:
        book.close()

    vintage = next(
        (_text(c) for row in rows[:12] for c in row if _text(c).lower().startswith("updated on")),
        None,
    )
    if vintage is None:
        raise SourceError("the sheet carries no 'Updated on' line, so the release cannot be dated")

    columns = {name: _find_column(rows, label) for name, label in SERIES.items()}
    first_data_row = max(r for r, _ in columns.values()) + 2

    records = []
    for row in rows[first_data_row:]:
        if not row:
            continue
        match = _PERIOD.match(_text(row[0]))
        if not match:
            continue
        when = pd.Timestamp(int(match.group(1)), int(match.group(2)), 1)
        if when < FIRST_MONTH:
            continue
        record: dict[str, Any] = {"date": when}
        for name, (_, column) in columns.items():
            raw = row[column] if column < len(row) else None
            record[name] = _number(raw, where="%s %s" % (SERIES[name], _text(row[0])))
        records.append(record)
    if not records:
        raise SourceError("no month from %s in the sheet" % FIRST_MONTH.strftime("%Y-%m"))
    frame = pd.DataFrame(records).sort_values("date", kind="mergesort").reset_index(drop=True)
    if frame["date"].duplicated().any():
        raise SourceError("the sheet repeats a month")
    return frame, vintage


def release_date(vintage: str) -> str:
    """'Updated on September 02, 2026' as 2026-09-02."""
    text = re.sub(r"(?i)^updated on\s+", "", vintage).strip()
    return pd.to_datetime(text, format="%B %d, %Y").strftime("%Y-%m-%d")


def compare_releases(before: pd.DataFrame, after: pd.DataFrame) -> pd.DataFrame:
    """Every value a later release changed, one row per month and series.

    A month the earlier release did not have is new data, not a revision.
    """
    rows = []
    old = before.set_index(pd.to_datetime(before["date"]))
    new = after.set_index(pd.to_datetime(after["date"]))
    for when in old.index.intersection(new.index):
        for column in SERIES:
            a, b = old.at[when, column], new.at[when, column]
            same = (pd.isna(a) and pd.isna(b)) or (not pd.isna(a) and not pd.isna(b) and a == b)
            if not same:
                rows.append(
                    {
                        "date": when,
                        "series": column,
                        "release_before": old.at[when, "release"],
                        "value_before": a,
                        "release_after": new.at[when, "release"],
                        "value_after": b,
                    }
                )
    frame = pd.DataFrame(
        rows, columns=["date", "series", "release_before", "value_before", "release_after", "value_after"]
    )
    return frame.sort_values(["date", "series"], kind="stable").reset_index(drop=True)


class WorldBankGasRevisions(Adapter):
    """Every value a Pink Sheet release changed, release against release."""

    name = "worldbank_gas_revisions"
    source = "World Bank Commodity Price Data (The Pink Sheet), compared release by release"
    url = LANDING_PAGE
    page_url = LANDING_PAGE
    unit = "USD per MMBtu"
    frequency = "monthly"
    method = "derived"
    unique_dates = False
    required_cols = ("date", "series", "release_before", "release_after")
    observation_column = "release_after"

    def __init__(self, additions: pd.DataFrame):
        self.additions = additions

    def fetch(self) -> pd.DataFrame:
        existing = read_cache(self.name, directory=self.directory())
        merged = pd.concat([f for f in (existing, self.additions) if f is not None], ignore_index=True)
        merged["date"] = pd.to_datetime(merged["date"])
        merged = merged.drop_duplicates(subset=["date", "series", "release_before", "release_after"])
        if len(self.additions):
            self.note = "%d value(s) changed by the release of %s against the release of %s." % (
                len(self.additions),
                self.additions["release_after"].iloc[0],
                self.additions["release_before"].iloc[0],
            )
        return merged.sort_values(["date", "series", "release_after"], kind="stable").reset_index(drop=True)


class WorldBankGasMonthly(Adapter):
    """Europe gas, US gas at Henry Hub and Japan LNG import price, monthly, from 2015."""

    name = "worldbank_gas_monthly"
    source = "World Bank Commodity Price Data (The Pink Sheet)"
    url = LANDING_PAGE
    page_url = LANDING_PAGE
    unit = "USD per MMBtu"
    frequency = "monthly"
    method = "published"
    required_cols = ("date", "release") + tuple(SERIES)
    bounds = {
        "europe_gas_usd_mmbtu": BOUNDS_LNG_USD_MMBTU,
        "henry_hub_usd_mmbtu": BOUNDS_HENRY_HUB_USD_MMBTU,
        "japan_lng_import_usd_mmbtu": BOUNDS_LNG_USD_MMBTU,
    }
    min_observations = {name: 100 for name in SERIES}
    observation_column = "europe_gas_usd_mmbtu"

    def __init__(self, *, from_file=None, fetched_at: str | None = None):
        self.from_file = from_file
        self.fetched_at = fetched_at
        self.revisions: pd.DataFrame | None = None

    def fetch(self) -> pd.DataFrame:
        if self.from_file is not None:
            payload = open(self.from_file, "rb").read()
            how = "read from a copy fetched at %s" % self.fetched_at
        else:
            workbook_url = discover_workbook_url(http_get(LANDING_PAGE).content)
            payload = http_get(workbook_url, timeout=60).content
            self.url = workbook_url
            how = "fetched from %s" % workbook_url
        frame, vintage = parse_workbook(payload)
        released = release_date(vintage)
        frame.insert(1, "release", released)

        existing = read_cache(self.name, directory=self.directory())
        if existing is not None and len(existing):
            committed = str(existing["release"].iloc[-1])
            if released < committed:
                raise SourceError(
                    "the workbook is the release of %s, older than the committed release of %s"
                    % (released, committed)
                )
            if released != committed:
                self.revisions = compare_releases(existing, frame)

        japan = frame[frame[PROVISIONAL_COLUMN].notna()]
        self.provisional_from = (
            japan["date"].iloc[-PROVISIONAL_MONTHS].strftime("%Y-%m-%d") if len(japan) >= PROVISIONAL_MONTHS else None
        )
        self.vintage = vintage
        self.note = (
            "%s. The Europe series is TTF from April 2015 and a border price before; "
            "the US series is the Henry Hub spot price; the Japan series is an LNG "
            "import price, cif, whose last two months the World Bank calls estimates, "
            "so provisional_from applies to that column only." % how
        )
        return frame

    def run(self) -> dict:
        entry = super().run()
        if self.revisions is not None and len(self.revisions):
            WorldBankGasRevisions(self.revisions).run()
        return entry


def main() -> int:
    try:
        entry = WorldBankGasMonthly().run()
    except Exception as exc:
        print("FAILED worldbank_gas_monthly: %s" % exc)
        return 1
    print("ok     %s: %d months, %s to %s, %s" % (
        entry["series"], entry["observations"], entry["first_date"], entry["last_date"], entry["vintage"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
