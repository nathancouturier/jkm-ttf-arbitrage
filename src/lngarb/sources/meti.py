"""Japan's spot LNG prices from METI's survey, monthly, USD/MMBtu, March 2014 to March 2021.

METI's Spot LNG Price Statistics are simple averages of fixed price spot cargoes
bought by companies that consume them in Japan, converted to a DES basis, in two
forms: contract-based, from March 2014, and arrival-based, from April 2014. METI
discontinued the survey with the figures for March 2021, released on 14 May
2021, and JOGMEC's survey continues it from April 2021. METI does not publish a
month when fewer than two reporters imported spot LNG; such a month is a gap,
never a zero.

Where it is read from. METI publishes a PDF for every month and one historical
workbook with every month's latest figure, labelled "Detailed", or "Fixed" for
the one month METI corrected again the following March. The workbook is the
source here: it holds the whole series, and it agrees with every monthly PDF
this study could read. METI's site answers automated requests with a bot
challenge after a handful of files, and the series is frozen, so nothing here
refetches it: the workbook fetched once is kept in data/private/meti/ and
committed as a test fixture, and the parsed series is committed. The monthly
PDFs, which alone carry the preliminary figures, are a manual step: those saved
by hand in data/private/meti/pdf/ are read by MetiSpotLngReleases, one row per
month a release prints, so that each figure is kept as first published and as
corrected (a vintage table).
"""

from __future__ import annotations

import io
import math
import re
import sys
from pathlib import Path
from typing import Any

import pandas as pd

from ..config import BOUNDS_LNG_USD_MMBTU
from . import base
from .base import Adapter, SourceError

__all__ = [
    "PAGE_URL", "WORKBOOK_URL", "parse_workbook", "MetiSpotLngMonthly", "workbook_path",
    "release_dir", "parse_release_text", "parse_release", "MetiSpotLngReleases",
]

PAGE_URL = "https://www.meti.go.jp/english/statistics/sho/slng/index.html"
WORKBOOK_URL = "https://www.meti.go.jp/english/statistics/sho/slng/historical-data-e.xlsx"
SHEET = "historical data"

#: METI marks a month it did not publish with a multiplication sign.
NOT_PUBLISHED = chr(0x00D7)


def workbook_path() -> Path:
    """Where the workbook fetched once from METI is kept, gitignored."""
    return base.PRIVATE / "meti" / "historical-data-e.xlsx"


def _number(cell: Any, *, where: str) -> float:
    if cell is None:
        raise SourceError("%s is empty; METI marks an unpublished month with %r" % (where, NOT_PUBLISHED))
    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
        return float(cell)
    if str(cell).strip() == NOT_PUBLISHED:
        return math.nan
    raise SourceError("%s holds %r, neither a price nor METI's gap marker" % (where, cell))


def parse_workbook(payload: bytes) -> pd.DataFrame:
    """Every month in METI's historical workbook, with METI's label for its figure."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        if SHEET not in book.sheetnames:
            raise SourceError("no %r sheet; sheets are %s" % (SHEET, book.sheetnames))
        rows = [list(r) for r in book[SHEET].iter_rows(values_only=True)]
    finally:
        book.close()

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if header[:2] != ["Year", "Month"] or header[3:5] != ["Contract-based", "Arrival-based"]:
        raise SourceError("the header row reads %r, not the layout this parser knows" % (header,))
    if "USD/MMBtu" not in header[5]:
        raise SourceError("the unit cell reads %r, expected USD/MMBtu" % (header[5],))

    records = []
    year = None
    for number, row in enumerate(rows[1:], start=2):
        if all(c is None for c in row):
            continue
        if row[0] is not None:
            year = int(row[0])
        if year is None or row[1] is None:
            raise SourceError("row %d has no year or month" % number)
        label = str(row[2]).strip() if row[2] is not None else ""
        if label not in ("Detailed", "Fixed"):
            raise SourceError("row %d labels its figure %r" % (number, label))
        records.append(
            {
                "date": pd.Timestamp(year, int(row[1]), 1),
                "contract_based_usd_mmbtu": _number(row[3], where="row %d contract-based" % number),
                "arrival_based_usd_mmbtu": _number(row[4], where="row %d arrival-based" % number),
                "figure": label,
            }
        )
    frame = pd.DataFrame(records)
    if not frame["date"].is_monotonic_increasing or frame["date"].duplicated().any():
        raise SourceError("the months are not in order or repeat")
    return frame


class MetiSpotLngMonthly(Adapter):
    """METI spot LNG prices, contract-based and arrival-based, March 2014 to March 2021."""

    name = "meti_spot_lng_monthly"
    source = "Ministry of Economy, Trade and Industry of Japan, Spot LNG Price Statistics"
    url = WORKBOOK_URL
    page_url = PAGE_URL
    unit = "USD per MMBtu, DES"
    frequency = "monthly"
    method = "published"
    required_cols = ("date", "contract_based_usd_mmbtu", "arrival_based_usd_mmbtu", "figure")
    bounds = {
        "contract_based_usd_mmbtu": BOUNDS_LNG_USD_MMBTU,
        "arrival_based_usd_mmbtu": BOUNDS_LNG_USD_MMBTU,
    }
    min_observations = {"contract_based_usd_mmbtu": 70, "arrival_based_usd_mmbtu": 70}
    observation_column = "contract_based_usd_mmbtu"

    def __init__(self, *, from_file: Path | None = None, fetched_at: str = "2026-09-30T16:56:44Z"):
        self.from_file = from_file
        self.fetched_at = fetched_at

    def fetch(self) -> pd.DataFrame:
        path = Path(self.from_file) if self.from_file is not None else workbook_path()
        if not path.exists():
            raise SourceError(
                "%s is not on this machine. METI's site challenges automated requests "
                "and the series is frozen, so the workbook is saved once by hand, not "
                "fetched by this pipeline" % path
            )
        frame = parse_workbook(path.read_bytes())
        gaps_c = int(frame["contract_based_usd_mmbtu"].isna().sum())
        gaps_a = int(frame["arrival_based_usd_mmbtu"].isna().sum())
        self.vintage = "historical workbook, series discontinued with March 2021"
        self.note = (
            "Read from METI's historical workbook, fetched from %s at %s and not "
            "refetched. %d months; METI did not publish %d contract-based and %d "
            "arrival-based months, when fewer than two reporters imported spot LNG. "
            "Created by processing the information in the Spot LNG Price Statistics "
            "(Ministry of Economy, Trade and Industry of Japan)."
            % (WORKBOOK_URL, self.fetched_at, len(frame), gaps_c, gaps_a)
        )
        return frame

    def _entry(self, *, status: str, frame, note: str) -> dict:
        # Read from a file fetched once, so the fetch time is that fetch's, not
        # this run's.
        entry = super()._entry(status=status, frame=frame, note=note)
        entry["fetched_at"] = self.fetched_at
        return entry


# --------------------------------------------------------------------------
# The monthly releases, saved by hand
# --------------------------------------------------------------------------

def release_dir() -> Path:
    """Where the monthly PDFs saved by hand are kept, gitignored, under the names METI links them by."""
    return base.PRIVATE / "meti" / "pdf"


_RELEASE_TITLE = re.compile(r"^\((Preliminary|Detailed) Figures for ([A-Z][a-z]+ \d{4})\)$")
_RELEASE_DAY = re.compile(r"^([A-Z][a-z]+ \d{1,2}, \d{4})$")
_RELEASE_ROW = re.compile(
    r"^(\d{4}) ([A-Z][a-z]+) ?(\*{1,3}) (\d+\.\d|" + NOT_PUBLISHED + r") (\d+\.\d|" + NOT_PUBLISHED + r")$"
)
_MARKS = {"*": "Detailed", "**": "Preliminary", "***": "Fixed"}


def parse_release_text(text: str, *, name: str) -> list[dict[str, Any]]:
    """The rows of one monthly release: each month it prints, with METI's label for its figure.

    A release prints the survey month as Preliminary, the month before as
    Detailed and, where METI corrected one again, a month of the year before as
    Fixed (the release for March 2020 does, that for March 2021 does not); the
    marks are read from the release's own key. The release's title, its date
    and its key must be found, or the release is refused.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    title = next((m for m in map(_RELEASE_TITLE.match, lines) if m), None)
    if title is None:
        raise SourceError("%s: no '(Preliminary Figures for ...)' title" % name)
    position = next(i for i, line in enumerate(lines) if _RELEASE_TITLE.match(line))
    day = _RELEASE_DAY.match(lines[position + 1]) if position + 1 < len(lines) else None
    if day is None:
        raise SourceError("%s: no release date under the title" % name)
    key = next((line for line in lines if line.startswith("*Detailed")), None)
    if key is None or (title.group(1) == "Preliminary" and "Preliminary" not in key):
        raise SourceError("%s: no key to its marks" % name)
    release_date = pd.Timestamp(pd.to_datetime(day.group(1), format="%B %d, %Y"))
    rows = []
    for line in lines:
        match = _RELEASE_ROW.match(line)
        if not match:
            continue
        year, month, mark, contract, arrival = match.groups()
        if mark == "***" and "Fixed" not in key:
            raise SourceError("%s: a row marked *** but no Fixed in the key" % name)
        rows.append({
            "date": pd.Timestamp(pd.to_datetime("%s %s" % (month, year), format="%B %Y")),
            "release_date": release_date,
            "release": name,
            "figure": _MARKS[mark],
            "contract_based_usd_mmbtu": math.nan if contract == NOT_PUBLISHED else float(contract),
            "arrival_based_usd_mmbtu": math.nan if arrival == NOT_PUBLISHED else float(arrival),
        })
    survey = pd.Timestamp(pd.to_datetime(title.group(2), format="%B %Y"))
    if title.group(1) == "Preliminary" and not any(r["date"] == survey and r["figure"] == "Preliminary" for r in rows):
        raise SourceError("%s: the table does not print the preliminary figure of %s" % (name, title.group(2)))
    if not rows:
        raise SourceError("%s: no table row read" % name)
    return rows


def parse_release(pdf: bytes, *, name: str) -> list[dict[str, Any]]:
    """The rows of one monthly release PDF (parse_release_text)."""
    import pdfplumber

    with pdfplumber.open(io.BytesIO(pdf)) as document:
        text = "\n".join((page.extract_text() or "") for page in document.pages)
    return parse_release_text(text, name=name)


class MetiSpotLngReleases(Adapter):
    """Every figure the monthly releases saved by hand print, preliminary, detailed and fixed, as first published."""

    name = "meti_spot_lng_releases"
    source = "Ministry of Economy, Trade and Industry of Japan, Spot LNG Price Statistics, monthly releases"
    url = PAGE_URL
    page_url = PAGE_URL
    unit = "USD per MMBtu, DES"
    frequency = "monthly"
    method = "parsed"
    required_cols = ("date", "release_date", "release", "figure", "contract_based_usd_mmbtu", "arrival_based_usd_mmbtu")
    bounds = {
        "contract_based_usd_mmbtu": BOUNDS_LNG_USD_MMBTU,
        "arrival_based_usd_mmbtu": BOUNDS_LNG_USD_MMBTU,
    }
    min_observations = {"contract_based_usd_mmbtu": 2, "arrival_based_usd_mmbtu": 2}
    observation_column = "contract_based_usd_mmbtu"
    #: a month is printed by two releases, or three, as it is revised
    unique_dates = False

    def __init__(self, *, releases: dict[str, bytes] | None = None):
        # releases maps a file name to its PDF, standing in for the folder
        self.releases = releases

    def fetch(self) -> pd.DataFrame:
        releases = self.releases
        if releases is None:
            folder = release_dir()
            paths = sorted(folder.glob("*.pdf")) if folder.exists() else []
            releases = {p.name: p.read_bytes() for p in paths}
            if paths:
                # Saved by hand: the fetch time is when the last file was saved.
                from datetime import datetime, timezone
                saved = max(p.stat().st_mtime for p in paths)
                self.fetched_at = datetime.fromtimestamp(saved, timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        if not releases:
            raise SourceError(
                "no monthly release is saved in %s; METI's site challenges automated "
                "requests, so the releases are saved by hand (the manual step)" % release_dir().name
            )
        rows = []
        for name, pdf in sorted(releases.items()):
            rows.extend(parse_release(pdf, name=name))
        frame = pd.DataFrame(rows).sort_values(["date", "release_date"]).reset_index(drop=True)
        frame["anomaly"] = None
        workbook = base.read_cache("meti_spot_lng_monthly")
        if workbook is not None:
            latest = workbook.set_index("date")
            for i, r in frame.iterrows():
                if r["figure"] == "Preliminary" or r["date"] not in latest.index:
                    continue
                kept = latest.loc[r["date"], "contract_based_usd_mmbtu"]
                if not (pd.isna(kept) and pd.isna(r["contract_based_usd_mmbtu"])) and kept != r["contract_based_usd_mmbtu"]:
                    frame.at[i, "anomaly"] = "the workbook holds %s for the contract-based figure" % kept
        months = frame["date"].dt.strftime("%B %Y")
        self.vintage = "%d releases saved by hand, %s to %s" % (
            len(releases), frame["release_date"].min().date(), frame["release_date"].max().date())
        self.note = (
            "Every figure printed by the %d monthly releases saved by hand, from %s to %s, "
            "each labelled as METI labels it (Preliminary, Detailed, Fixed) with the day "
            "the release was published, so that a figure is kept as first published and "
            "as corrected. Created by processing the information in the Spot LNG Price "
            "Statistics (Ministry of Economy, Trade and Industry of Japan)."
            % (len(releases), months.iloc[0], months.iloc[-1])
        )
        return frame


def main() -> int:
    try:
        entry = MetiSpotLngMonthly().run()
    except Exception as exc:
        print("FAILED meti_spot_lng_monthly: %s" % exc)
        return 1
    print("ok     %s: %d months, %s to %s" % (entry["series"], entry["observations"], entry["first_date"], entry["last_date"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
