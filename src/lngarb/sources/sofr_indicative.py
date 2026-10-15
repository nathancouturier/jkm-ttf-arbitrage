"""The indicative SOFR the New York Fed computed for August 2014 to March 2018, before it published SOFR.

SOFR's first value date is 2 April 2018. For the years before it the New York
Fed published, once, a workbook of indicative rates computed by the same
method over the same transactions, 22 August 2014 to 29 March 2018: the
tri-party general collateral rate, the broad general collateral rate and the
Secured Overnight Financing Rate, as volume weighted medians in basis points,
with their volumes. The study reads the SOFR column, so that the financing
line of a cargo loading from August 2014 runs on a secured rate like the
cargoes after April 2018, and the effective federal funds rate, unsecured,
stands in only before August 2014.

The workbook is "Data Release.xlsx", linked from the New York Fed's page of
additional information about its reference rates as "Daily indicative TGCR,
BGCR, and SOFR volume and rate data (Aug. 2014 - Mar. 2018)". It is a closed
file: nothing is added to it. Its sheet "VWM Rates" holds one row per
business day; the rates are in basis points and are read into percent. The
New York Fed's Terms of Use cover it, with the reference rate notice.
"""

from __future__ import annotations

import io
import sys

import pandas as pd

from ..config import BOUNDS_SOFR_PERCENT
from .base import Adapter, SourceError, http_get

__all__ = ["FILE_URL", "PAGE_URL", "FIRST_VALUE_DATE", "LAST_VALUE_DATE", "parse_workbook", "SofrIndicativeDaily"]

FILE_URL = "https://www.newyorkfed.org/medialibrary/media/markets/Data%20Release.xlsx"
PAGE_URL = "https://www.newyorkfed.org/markets/reference-rates/additional-information-about-reference-rates"
SHEET = "VWM Rates"
COLUMN = "Secured Overnight Financing Rate"
FIRST_VALUE_DATE = "2014-08-22"
LAST_VALUE_DATE = "2018-03-29"
#: basis points in a percentage point
BP_PER_PERCENT = 100.0


def parse_workbook(payload: bytes) -> pd.DataFrame:
    """The indicative SOFR, daily, percent per year, from the workbook's rates sheet.

    Raises SourceError when the sheet or the column is missing, the units cell
    does not say basis points, a rate is not a number, or the dates are not
    the closed window the New York Fed published.
    """
    import openpyxl

    try:
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001, openpyxl raises several types
        raise SourceError("the download did not open as an xlsx workbook, %s: %s" % (type(exc).__name__, exc)) from exc
    try:
        if SHEET not in book.sheetnames:
            raise SourceError("no %r sheet; sheets are %s" % (SHEET, ", ".join(book.sheetnames)))
        rows = [list(row) for row in book[SHEET].iter_rows(values_only=True)]
    finally:
        book.close()
    header = next((i for i, row in enumerate(rows[:5]) if row and row[0] == "Date"), None)
    if header is None:
        raise SourceError("the rates sheet has no Date header in its first rows")
    names = [str(c).strip() if c is not None else "" for c in rows[header]]
    if COLUMN not in names:
        raise SourceError("the rates sheet has no %r column; columns are %s" % (COLUMN, names))
    column = names.index(COLUMN)
    # The legend beside the table names the series and its units.
    legend = [str(c) for row in rows[header:header + 6] for c in row[column + 1:] if c is not None]
    if not any("Basis points" in c for c in legend):
        raise SourceError("the rates sheet's legend does not say the rates are in basis points")
    records = []
    for row in rows[header + 1:]:
        if not row or row[0] is None:
            continue
        day = pd.Timestamp(row[0])
        value = row[column]
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            raise SourceError("the indicative SOFR on %s is %r, not a number" % (day.date(), value))
        records.append({"date": day.normalize(), "sofr_percent": float(value) / BP_PER_PERCENT})
    frame = pd.DataFrame(records).sort_values("date", kind="mergesort").reset_index(drop=True)
    if frame.empty or frame["date"].duplicated().any():
        raise SourceError("the rates sheet holds no day, or repeats one")
    first, last = frame["date"].iloc[0], frame["date"].iloc[-1]
    if first != pd.Timestamp(FIRST_VALUE_DATE) or last != pd.Timestamp(LAST_VALUE_DATE):
        raise SourceError("the indicative rates run from %s to %s, not %s to %s"
                          % (first.date(), last.date(), FIRST_VALUE_DATE, LAST_VALUE_DATE))
    return frame


class SofrIndicativeDaily(Adapter):
    """The indicative SOFR, daily, percent, 22 August 2014 to 29 March 2018, from the New York Fed."""

    name = "nyfed_sofr_indicative_daily"
    source = "Federal Reserve Bank of New York, indicative Secured Overnight Financing Rate, August 2014 to March 2018"
    url = FILE_URL
    page_url = PAGE_URL
    unit = "percent per year"
    frequency = "daily"
    method = "published"
    required_cols = ("date", "sofr_percent")
    bounds = {"sofr_percent": BOUNDS_SOFR_PERCENT}
    min_observations = {"sofr_percent": 850}
    observation_column = "sofr_percent"

    def __init__(self, *, from_file=None, fetched_at: str | None = None):
        self.from_file = from_file
        self.fetched_at = fetched_at

    def fetch(self) -> pd.DataFrame:
        if self.from_file is not None:
            payload = open(self.from_file, "rb").read()
            how = "read from a copy fetched at %s" % self.fetched_at
        else:
            payload = http_get(FILE_URL, timeout=120).content
            how = "fetched from the New York Fed's workbook of indicative rates"
        frame = parse_workbook(payload)
        self.vintage = "value dates %s to %s, a closed file" % (FIRST_VALUE_DATE, LAST_VALUE_DATE)
        self.note = (
            "%s, sheet %r, column %r, basis points read into percent, %s to %s. The New York "
            "Fed computed these rates by SOFR's method over the same transactions before it "
            "published SOFR, whose first value date is 2 April 2018; the financing line runs on "
            "them for a cargo loading in that span, a secured rate like SOFR, and on the effective "
            "federal funds rate, unsecured, before 22 August 2014. The Secured Overnight Financing "
            "Rate is subject to the Terms of Use posted at newyorkfed.org; the New York Fed is not "
            "responsible for its publication here, does not sanction or endorse any particular "
            "republication, and has no liability for its use."
            % (how, SHEET, COLUMN, FIRST_VALUE_DATE, LAST_VALUE_DATE)
        )
        return frame


def main() -> int:
    try:
        entry = SofrIndicativeDaily().run()
    except Exception as exc:
        print("FAILED nyfed_sofr_indicative_daily: %s" % exc)
        return 1
    print("ok     %s: %d observations, %s to %s" % (
        entry["series"], entry["observations"], entry["first_date"], entry["last_date"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
